"""Regression checks for identity, translation safety and portable HTML generation."""
import copy
import importlib.util
import json
import subprocess
import sys
import unittest
import contextlib
import shutil
import uuid
from pathlib import Path

HOME = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(HOME))
from common import read, write, load_config, validate_translations
from diff_blocks import compare
from prepare_translations import prepare
from build_review_page import script_json

@contextlib.contextmanager
def scratch():
    # Use a fixture-local directory so restricted runners need no system-temp access.
    parent = (HOME / 'examples/output/test-tmp').resolve()
    parent.mkdir(parents=True, exist_ok=True)
    directory = parent / uuid.uuid4().hex
    directory.mkdir()
    try:
        yield str(directory)
    finally:
        if directory.resolve().parent != parent:
            raise RuntimeError('Refusing cleanup outside the fixture temporary directory')
        shutil.rmtree(directory)

def block(bid, text, kind='text'):
    return {'id': bid, 'doc': 'main', 'kind': kind, 'en': text, 'match': text, 'html': text}

def data(blocks):
    return {'docs': {'main': {}}, 'blocks': blocks}

class PipelineTests(unittest.TestCase):
    def test_insert_delete_and_repeated_text(self):
        old = data([block('O1','alpha identical text'), block('O2','repeated words'), block('O3','repeated words'), block('O4','removed entirely', 'reference')])
        new = data([block('N1','alpha identical text'), block('N2','brand new paragraph about stars'), block('N3','repeated words'), block('N4','repeated words')])
        result = compare(old, new, HOME, HOME)
        self.assertEqual(result['N2']['status'], 'new')
        self.assertEqual([result[k]['prev'] for k in ('N1','N3','N4')], ['O1','O2','O3'])
        self.assertEqual([b['id'] for b in result['_dropped_main']], ['O4'])

    def test_last_table_cell_change(self):
        table = {'id':'O', 'kind':'table', 'doc':'main', 'rows':[[{'text':'Shelf'}, {'text':'Books'}],[{'text':'Cloud'},{'text':'6'}]]}
        newer = copy.deepcopy(table)
        newer['id'] = 'N'
        newer['rows'][-1][-1]['text'] = '9'
        self.assertEqual(compare(data([table]), data([newer]), HOME, HOME)['N']['status'], 'changed')

    def test_changes_are_monotone(self):
        old = data([block('O1','The red fox sleeps in a cave.'), block('O2','The blue bird sleeps in a tree.')])
        new = data([block('N1','The blue bird rests in a tree.'), block('N2','The red fox rests in a cave.')])
        out = compare(old,new,HOME,HOME)
        ids = [out[b['id']]['prev'] for b in new['blocks'] if out[b['id']]['prev']]
        self.assertEqual(ids, sorted(ids))

    def test_translation_coverage_rejects_missing_and_extra(self):
        blocks = [block('A','A paragraph')]
        for strict, summary in [({}, {'A':'摘要'}), ({'A':''},{'A':'摘要'}), ({'A':'翻译','Z':'错误'},{'A':'摘要'})]:
            with self.assertRaises(ValueError): validate_translations(blocks, strict, summary)

    def test_example_reuse_and_reference_context(self):
        c = load_config(HOME / 'examples/config_v2.json')
        work = prepare(c)
        self.assertEqual(len(work['pending']), 4)
        self.assertIn('M-006', work['strict'])
        changed = next(p for p in work['pending'] if p['id'] == 'M-003')
        self.assertIn('Friday', changed['previous_source'])
        self.assertIn('周五', changed['previous_strict'])

    def test_stale_diff_cannot_reuse_changed_translation(self):
        with scratch() as td:
            tmp = Path(td)
            c = {'_base':tmp, 'blocks':'new.json','previous_blocks':'old.json','diff':'diff.json','previous_strict':'strict.json','previous_summary':'summary.json'}
            write(tmp/'old.json',data([block('O','old words')]))
            write(tmp/'new.json',data([block('N','new words')]))
            write(tmp/'diff.json',{'N':{'prev':'O','status':'same'}})
            write(tmp/'strict.json',{'O':'旧译文'})
            write(tmp/'summary.json',{'O':'旧摘要'})
            work = prepare(c)
            self.assertFalse(work['strict'])
            self.assertEqual(work['pending'][0]['needed'], ['strict','summary'])

    def test_script_metadata_escaping(self):
        value = {'title':'</script><script>alert(1)</script>'}
        encoded = script_json(value)
        self.assertNotIn('<', encoded)
        self.assertEqual(json.loads(encoded), value)

    def test_optional_pdf_unique_page_and_ambiguous_text(self):
        try:
            import fitz
        except ImportError:
            self.skipTest('Optional PDF package is not installed')
        from docx import Document
        from extract_blocks import extract
        import word_helpers
        with scratch() as td:
            tmp = Path(td)
            doc = Document()
            doc.add_paragraph('A unique reading bench sits beside the window.')
            doc.add_paragraph('Repeated room label')
            doc.save(tmp/'story.docx')
            pdf = fitz.open()
            pdf.new_page().insert_text((50,50),'A unique reading bench sits beside the window.\nRepeated room label')
            pdf.new_page().insert_text((50,50),'Repeated room label')
            pdf.save(tmp/'story.pdf')
            pdf.close()
            word_helpers.ASSETS = tmp/'assets'
            word_helpers.ASSETS.mkdir()
            blocks, meta = extract('main',{'docx':'story.docx','pdf':'story.pdf'},tmp,tmp/'blocks.json')
            self.assertEqual(blocks[0]['page'],1)
            self.assertIsNone(blocks[1]['page'])
            self.assertEqual(meta['pdf_pages'],2)

    def test_main_only_without_diff(self):
        with scratch() as td:
            tmp = Path(td)
            c = read(HOME/'examples/config_v1.json')
            base = HOME/'examples'
            c['documents'] = {'main':c['documents']['main']}
            original = read(base/c['blocks'])
            original['docs'].pop('si')
            original['blocks'] = [b for b in original['blocks'] if b['doc']=='main']
            # Keep assets adjacent to the temporary blocks JSON.
            import shutil
            shutil.copytree(base/'output/v1/assets',tmp/'assets')
            for field in ('strict','summary'):
                write(tmp/f'{field}.json',{k:v for k,v in read(base/c[field]).items() if k.startswith('M-')})
                c[field] = f'{field}.json'
            write(tmp/'blocks.json',original)
            c.update(blocks='blocks.json', html='review.html', page_notes=None, title='Demo </script> title', languages={'source':'Français','target':'日本語'})
            write(tmp/'config.json',c)
            subprocess.run([sys.executable,'-B',str(HOME/'build_review_page.py'),'--config',str(tmp/'config.json')],check=True,capture_output=True)
            page = (tmp/'review.html').read_text(encoding='utf-8')
            self.assertIn('Français', page)
            self.assertNotIn('data-doc="si"', page)
            self.assertEqual(page.count('</script>'),1)
            self.assertIn('data:image/png;base64,',page)

    def test_private_leak_scan_zip_and_embedded_image(self):
        spec = importlib.util.spec_from_file_location('scanner',HOME/'tools/leak_scan.py')
        scanner = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(scanner)
        import base64
        import zipfile
        with scratch() as td:
            tmp=Path(td)
            with zipfile.ZipFile(tmp/'fixture.docx','w') as archive: archive.writestr('word/document.xml','secret test token')
            encoded=base64.b64encode(b'\x00secret test token').decode()
            (tmp/'page.html').write_text('data:image/png;base64,'+encoded)
            out=scanner.scan(tmp,[{'id':'fixture','pattern':'secret test token'}])
            self.assertEqual(out['status'],'failed')
            self.assertTrue(any('!word/document.xml' in f['file'] for f in out['findings']))
            self.assertTrue(any('!image-' in f['file'] for f in out['findings']))

if __name__ == '__main__': unittest.main()
